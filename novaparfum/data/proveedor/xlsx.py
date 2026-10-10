import json, csv
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
items=json.load(open('items.json'))
HDR=['N°','Categoría','Género','Marca','Perfume','Concentración','Tamaño (ml)','¿Lo tenés?','Precio por unidad ($U)','Observaciones','Estado en NovaParfum']
wb=Workbook(); ws=wb.active; ws.title='Lista para cotizar'
F='Arial'
ws['A1']='Lista de perfumes para cotizar · NovaParfum'; ws['A1'].font=Font(name=F,bold=True,size=14,color='2F2A4A')
ws['A2']='Completá las columnas amarillas en cada tamaño: si lo tenés (Sí/No), el precio por unidad en pesos uruguayos y cualquier aclaración (otro tamaño, tester, versión distinta).'
ws['A3']='Si un perfume lo tenés en un tamaño que no está en la lista, anotalo en Observaciones. Si no lo trabajás, poné "No".'
for c in ('A2','A3'): ws[c].font=Font(name=F,size=10,color='555555')
ws.append([]); ws.append(HDR)
hr=5
ink=PatternFill('solid',fgColor='2F2A4A'); inp=PatternFill('solid',fgColor='FFF3C4'); band=PatternFill('solid',fgColor='F3F0FA')
thin=Side(style='thin',color='D9D4E7')
for c in ws[hr]: c.font=Font(name=F,bold=True,color='F7F4F0'); c.fill=ink; c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
ws.row_dimensions[hr].height=32
# fila de ejemplo
ex=['Ejemplo','Árabe','Dama','Lattafa','Yara','EDP',100,'Sí',2050,'También tengo Yara Moi y Yara Tous','']
ws.append(ex)
for c in ws[hr+1]: c.font=Font(name=F,italic=True,color='8A8A8A')
r=hr+2; n=0
csvrows=[HDR,ex]
for it in items:
    n+=1
    for s in it['sizes']:
        row=[n,it['cat'],it['gen'],it['marca'],it['nombre'],it['conc'],s,None,None,None,it['estado']]
        ws.append(row); csvrows.append([x if x is not None else '' for x in row])
        for c in ws[r]:
            c.font=Font(name=F,size=10); c.border=Border(bottom=thin)
            if n%2==0: c.fill=band
        for col in (8,9,10): ws.cell(r,col).fill=inp
        ws.cell(r,9).number_format='#,##0'
        r+=1
last=r-1
dv=DataValidation(type='list',formula1='"Sí,No"',allow_blank=True); ws.add_data_validation(dv); dv.add(f'H{hr+2}:H{last}')
dv2=DataValidation(type='decimal',operator='greaterThan',formula1='0',allow_blank=True,error='Poné el precio en pesos, sin puntos ni símbolo.'); ws.add_data_validation(dv2); dv2.add(f'I{hr+2}:I{last}')
for i,w in enumerate([6,11,11,22,38,13,11,11,16,40,20],1): ws.column_dimensions[get_column_letter(i)].width=w
for rr in ws.iter_rows(min_row=hr+1,max_row=last):
    for idx in (0,6,7): rr[idx].alignment=Alignment(horizontal='center')
ws.freeze_panes=f'F{hr+1}'; ws.auto_filter.ref=f'A{hr}:K{last}'
# Resumen con fórmulas
rs=wb.create_sheet('Resumen')
rs['A1']='Resumen de la cotización'; rs['A1'].font=Font(name=F,bold=True,size=13,color='2F2A4A')
rs.append([])
rs.append(['Categoría','Filas (perfume y tamaño)','Lo tiene (Sí)','No lo tiene','Con precio cargado'])
for c in rs[3]: c.font=Font(name=F,bold=True,color='F7F4F0'); c.fill=ink; c.alignment=Alignment(wrap_text=True,horizontal='center')
L="'Lista para cotizar'"
for i,cat in enumerate(['Árabe','Diseñador','Nicho']):
    rr=4+i
    rs.append([cat,f'=COUNTIFS({L}!$B${hr+2}:$B${last},A{rr})',f'=COUNTIFS({L}!$B${hr+2}:$B${last},A{rr},{L}!$H${hr+2}:$H${last},"Sí")',
               f'=COUNTIFS({L}!$B${hr+2}:$B${last},A{rr},{L}!$H${hr+2}:$H${last},"No")',f'=COUNTIFS({L}!$B${hr+2}:$B${last},A{rr},{L}!$I${hr+2}:$I${last},">0")'])
rs.append(['Total','=SUM(B4:B6)','=SUM(C4:C6)','=SUM(D4:D6)','=SUM(E4:E6)'])
for c in rs[7]: c.font=Font(name=F,bold=True)
for row in rs.iter_rows(min_row=4,max_row=6):
    for c in row: c.font=Font(name=F)
for i,w in enumerate([14,22,14,14,18],1): rs.column_dimensions[get_column_letter(i)].width=w
wb.save('NovaParfum-lista-para-cotizar.xlsx')
with open('NovaParfum-lista-para-cotizar.csv','w',newline='',encoding='utf-8') as f: csv.writer(f).writerows(csvrows)
print('filas datos', last-hr-1, 'perfumes', n)
