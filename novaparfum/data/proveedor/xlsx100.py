import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
items=json.load(open('items.json'))
con=[i for i in items if 100 in i['sizes']]
sin=[dict(i, size=min(i['sizes'], key=lambda s:(abs(s-100), -s))) for i in items if 100 not in i['sizes']]
F='Arial'; ink=PatternFill('solid',fgColor='2F2A4A'); inp=PatternFill('solid',fgColor='FFF3C4'); band=PatternFill('solid',fgColor='F3F0FA'); thin=Side(style='thin',color='D9D4E7')
HDR=['N°','Categoría','Género','Marca','Perfume','Concentración','Tamaño (ml)','¿Lo tenés?','Precio por unidad ($U)','Observaciones']
wb=Workbook()
def hoja(ws, titulo, lineas, data, sizekey):
    ws['A1']=titulo; ws['A1'].font=Font(name=F,bold=True,size=14,color='2F2A4A')
    for k,l in enumerate(lineas): 
        c=ws.cell(2+k,1,l); c.font=Font(name=F,size=10,color='555555')
    hr=5; ws.append([]) if ws.max_row<4 else None
    for j,h in enumerate(HDR,1):
        c=ws.cell(hr,j,h); c.font=Font(name=F,bold=True,color='F7F4F0'); c.fill=ink; c.alignment=Alignment(wrap_text=True,vertical='center',horizontal='center')
    ws.row_dimensions[hr].height=32
    ex=['Ejemplo','Árabe','Dama','Lattafa','Yara','EDP',100,'Sí',2050,'También tengo Yara Moi y Yara Tous']
    for j,v in enumerate(ex,1): c=ws.cell(hr+1,j,v); c.font=Font(name=F,italic=True,color='8A8A8A')
    r=hr+2
    for n,it in enumerate(data,1):
        row=[n,it['cat'],it['gen'],it['marca'],it['nombre'],it['conc'],sizekey(it),None,None,None]
        for j,v in enumerate(row,1):
            c=ws.cell(r,j,v); c.font=Font(name=F,size=10); c.border=Border(bottom=thin)
            if n%2==0: c.fill=band
        for col in (8,9,10): ws.cell(r,col).fill=inp
        ws.cell(r,9).number_format='#,##0'
        for col in (1,7,8): ws.cell(r,col).alignment=Alignment(horizontal='center')
        r+=1
    last=r-1
    dv=DataValidation(type='list',formula1='"Sí,No"',allow_blank=True); ws.add_data_validation(dv); dv.add(f'H{hr+2}:H{last}')
    dv2=DataValidation(type='decimal',operator='greaterThan',formula1='0',allow_blank=True,error='Poné el precio en pesos, sin puntos ni símbolo.'); ws.add_data_validation(dv2); dv2.add(f'I{hr+2}:I{last}')
    for i,w in enumerate([6,11,11,22,38,13,11,11,16,44],1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes=f'F{hr+1}'; ws.auto_filter.ref=f'A{hr}:J{last}'
    return hr+2,last
ws1=wb.active; ws1.title='Lista 100 ml'
a1,z1=hoja(ws1,'Lista de perfumes para cotizar · NovaParfum · 100 ml',
  ['Todos en presentación de 100 ml. Completá las columnas amarillas: si lo tenés (Sí/No), el precio por unidad en pesos uruguayos y cualquier aclaración.',
   'Si no lo trabajás, poné "No". Si lo tenés en otra presentación (tester, set, otra concentración), anotalo en Observaciones.'], con, lambda i:100)
ws2=wb.create_sheet('Sin versión de 100 ml')
a2,z2=hoja(ws2,'Perfumes que no se fabrican en 100 ml',
  ['Estos perfumes salen en otros tamaños. Va el tamaño más cercano a 100 ml.',
   'Si preferís no cotizarlos, podés dejar esta hoja sin completar.'], sin, lambda i:i['size'])
rs=wb.create_sheet('Resumen')
rs['A1']='Resumen de la cotización'; rs['A1'].font=Font(name=F,bold=True,size=13,color='2F2A4A')
hdr=['Hoja','Categoría','Perfumes','Lo tiene (Sí)','No lo tiene','Con precio cargado']
for j,h in enumerate(hdr,1):
    c=rs.cell(3,j,h); c.font=Font(name=F,bold=True,color='F7F4F0'); c.fill=ink; c.alignment=Alignment(wrap_text=True,horizontal='center')
r=4
for sh,a,z in (("'Lista 100 ml'",a1,z1),("'Sin versión de 100 ml'",a2,z2)):
    for cat in ('Árabe','Diseñador','Nicho'):
        rs.cell(r,1,sh.strip("'")); rs.cell(r,2,cat)
        rs.cell(r,3,f'=COUNTIFS({sh}!$B${a}:$B${z},B{r})')
        rs.cell(r,4,f'=COUNTIFS({sh}!$B${a}:$B${z},B{r},{sh}!$H${a}:$H${z},"Sí")')
        rs.cell(r,5,f'=COUNTIFS({sh}!$B${a}:$B${z},B{r},{sh}!$H${a}:$H${z},"No")')
        rs.cell(r,6,f'=COUNTIFS({sh}!$B${a}:$B${z},B{r},{sh}!$I${a}:$I${z},">0")')
        for j in range(1,7): rs.cell(r,j).font=Font(name=F)
        r+=1
rs.cell(r,1,'Total').font=Font(name=F,bold=True)
for j,col in enumerate('CDEF',3): c=rs.cell(r,j,f'=SUM({col}4:{col}{r-1})'); c.font=Font(name=F,bold=True)
for i,w in enumerate([22,12,12,13,13,17],1): rs.column_dimensions[get_column_letter(i)].width=w
wb.save('NovaParfum-lista-100ml.xlsx')
print('con 100 ml:',len(con),' sin 100 ml:',len(sin))
print('sin 100:', '; '.join(f"{i['marca']} {i['nombre']} ({i['size']})" for i in sin))
